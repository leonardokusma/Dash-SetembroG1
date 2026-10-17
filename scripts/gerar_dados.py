"""Gera os dados do painel ALP a partir da planilha de faturamento.

Uso:
    python scripts/gerar_dados.py "C:/caminho/FaturamentoG1ALP.xlsx" --mes "Setembro 2026"

Lê as abas "Base" e "Faturamento por Placa" da planilha, junta com
dados/sienge.json (custos do Sienge, preenchido à mão a partir do PDF)
e grava:
    dados/viagens.json
    dados/faturamento_placa.json
    dados/dados.js   <- é este que o index.html carrega
"""
import argparse
import datetime
import json
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
DADOS = RAIZ / "dados"

# Grafias diferentes do mesmo destino na base
DESTINOS = {"JAGUARIAIVA": "Jaguariaiva", "CHAPECO": "Chapeco", "LONDRINA": "Londrina", "TOCO": "Toco"}


def normaliza(texto):
    s = str(texto).strip()
    return DESTINOS.get(s.upper(), s)


def le_viagens(wb):
    ws = wb["Base"]
    viagens = []
    for linha in ws.iter_rows(min_row=3, values_only=True):
        _nf, data, motorista, cte, cavalo, carreta, origem, destino, qtd, valor = linha[:10]
        if not isinstance(data, datetime.datetime):
            continue
        kg = qtd if isinstance(qtd, (int, float)) else None
        viagens.append({
            "data": data.strftime("%Y-%m-%d"),
            "motorista": str(motorista).strip(),
            "cte": str(cte),
            "cavalo": str(cavalo).strip(),
            "carreta": str(carreta).strip(),
            "origem": normaliza(origem),
            "destino": normaliza(destino),
            "quantidade_kg": kg,
            "quantidade_original": None if kg is not None or qtd is None else str(qtd),
            "valor_frete": round(float(valor), 2),
        })
    viagens.sort(key=lambda v: (v["data"], v["cavalo"], v["cte"]))
    for i, v in enumerate(viagens):
        v["id"] = f"V{i + 1:03d}"
    return viagens


def le_faturamento(wb):
    ws = wb["Faturamento por Placa"]
    linhas = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[1] and r[0] != "TOTAL" and isinstance(r[2], (int, float)):
            f = {
                "placa": str(r[1]).strip(),
                "fretes": int(r[2]),
                "peso_kg": r[3],
                "faturamento": round(r[4], 2),
                "repasse_alp": round(r[5], 2),
                "combustivel": round(r[6], 2),
                "parcela": round(r[7], 2),
                "manutencao_ipva_seguro": round(r[8], 2),
                "livre_g1": round(r[9], 2),
            }
            # Diferença entre o Livre G1 da planilha e a conta refeita (ex.: "+10000" na fórmula)
            conta = f["faturamento"] - f["repasse_alp"] - f["combustivel"] - f["parcela"] - f["manutencao_ipva_seguro"]
            f["ajuste_planilha"] = round(f["livre_g1"] - conta, 2) or 0.0
            linhas.append(f)
    return linhas


def confere(viagens, faturamento, sienge):
    avisos = []
    for f in faturamento:
        vs = [v for v in viagens if v["cavalo"] == f["placa"]]
        soma = round(sum(v["valor_frete"] for v in vs), 2)
        if len(vs) != f["fretes"] or abs(soma - f["faturamento"]) > 0.01:
            avisos.append(f"{f['placa']}: base tem {len(vs)} fretes / R$ {soma}, aba por placa tem {f['fretes']} / R$ {f['faturamento']}")
        if not any(s["placa"] == f["placa"] for s in sienge):
            avisos.append(f"{f['placa']}: sem centro de custo em dados/sienge.json")
    return avisos


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("planilha", help="caminho da planilha FaturamentoG1ALP (.xlsx)")
    ap.add_argument("--mes", default="Setembro 2026", help='mês de referência mostrado no painel, ex.: "Outubro 2026"')
    ap.add_argument("--periodo-sienge", default="22/09/2026 a 07/10/2026", help="período do relatório do Sienge")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.planilha, data_only=True)
    viagens = le_viagens(wb)
    faturamento = le_faturamento(wb)
    sienge = json.loads((DADOS / "sienge.json").read_text(encoding="utf-8"))

    for aviso in confere(viagens, faturamento, sienge):
        print("AVISO:", aviso, file=sys.stderr)

    (DADOS / "viagens.json").write_text(json.dumps(viagens, ensure_ascii=False, indent=1), encoding="utf-8")
    (DADOS / "faturamento_placa.json").write_text(json.dumps(faturamento, ensure_ascii=False, indent=1), encoding="utf-8")

    pacote = {
        "mes": args.mes,
        "periodo_sienge": args.periodo_sienge,
        "planilha": Path(args.planilha).name,
        "gerado_em": datetime.datetime.now().strftime("%d/%m/%Y %H:%M"),
        "viagens": viagens,
        "faturamento_placa": faturamento,
        "sienge": sienge,
    }
    js = "// Gerado por scripts/gerar_dados.py. Não edite à mão.\nwindow.DADOS = " + json.dumps(pacote, ensure_ascii=False) + ";\n"
    (DADOS / "dados.js").write_text(js, encoding="utf-8")
    total = sum(v["valor_frete"] for v in viagens)
    print(f"OK: {len(viagens)} fretes, R$ {total:,.2f}, {len(faturamento)} placas, {len(sienge)} linhas do Sienge -> dados/dados.js")


if __name__ == "__main__":
    main()
